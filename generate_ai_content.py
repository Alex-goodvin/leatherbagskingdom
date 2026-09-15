import base64
import json
import hashlib
import os
import re
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook


# =========================================================
# generate_ai_content.py
#
# goods(4).xlsx
#       ↓
# ACTIVE products
#       ↓
# AI inputs + image
#       ↓
# SHA-256
#       ↓
# existing JSON?
#       ↓
# hash same?
#   YES         NO
#    ↓           ↓
#   SKIP       AI API
#                ↓
#            new JSON
#
# update_site.py does NOT call AI.
# =========================================================


# ---------------------------------------------------------
# 1. Settings
# ---------------------------------------------------------

EXCEL_PATH = Path("goods(4).xlsx")

AI_RESULTS_DIR = Path("ai_results")

PRODUCT_SOURCES_DIR = Path("product_sources")


# ---------------------------------------------------------
# 2. Helpers
# ---------------------------------------------------------

def clean_text(value: Any) -> str:
    """Convert a value to clean text."""
    if value is None:
        return ""

    return str(value).strip()


def extract_listing_id(etsy_url: Any) -> str | None:
    """Extract listing ID from Etsy URL."""

    if not etsy_url:
        return None

    etsy_url = str(etsy_url).strip()

    match = re.search(
        r"/listing/(\d+)",
        etsy_url
    )

    if match:
        return match.group(1)

    return None


# ---------------------------------------------------------
# 3. Calculate source hash
# ---------------------------------------------------------

def calculate_source_hash(
    product_name: str,
    category: str,
    short_description: str,
    alt_text: str,
    original_description: str,
    image_path: Path
) -> str:

    hasher = hashlib.sha256()

    text_inputs = [
        ("product_name", product_name),
        ("category", category),
        ("short_description", short_description),
        ("alt_text", alt_text),
        ("original_description", original_description),
    ]

    for field_name, value in text_inputs:

        field_name_bytes = (
            field_name.encode("utf-8")
        )

        value_bytes = (
            value.encode("utf-8")
        )

        hasher.update(
            len(field_name_bytes).to_bytes(
                4,
                "big"
            )
        )

        hasher.update(
            field_name_bytes
        )

        hasher.update(
            len(value_bytes).to_bytes(
                8,
                "big"
            )
        )

        hasher.update(
            value_bytes
        )

    image_bytes = (
        image_path.read_bytes()
    )

    hasher.update(
        b"__IMAGE_BYTES__"
    )

    hasher.update(
        image_bytes
    )

    return hasher.hexdigest()


# ---------------------------------------------------------
# 4. Load cached result
# ---------------------------------------------------------

def load_cached_result(
    output_path: Path
) -> dict | None:

    if not output_path.exists():
        return None

    try:

        return json.loads(
            output_path.read_text(
                encoding="utf-8"
            )
        )

    except (
        OSError,
        json.JSONDecodeError
    ):

        return None


# ---------------------------------------------------------
# 5. Build AI prompt
# ---------------------------------------------------------

def build_prompt(
    product: dict,
    etsy_description: str
) -> str:

    return f"""
You are an expert product copywriter and SEO editor for a handmade
leather goods website.

Create website content for the following product.

PRODUCT NAME:
{product["name"]}

CATEGORY:
{product["category"]}

SHORT DESCRIPTION:
{product["short_description"]}

ORIGINAL ETSY DESCRIPTION:
{etsy_description}

ALT TEXT FROM EXCEL:
{product["alt_text"]}

IMPORTANT RULES:

1. The original Etsy description is the primary source of factual
   product information.

2. Use the product image as a visual source.

3. Only describe visual details that are clearly visible.

4. Do not invent dimensions, materials, features, pockets,
   hardware, closures or other specifications.

5. Do not turn an uncertain visual interpretation into a factual
   product specification.

6. Color is especially sensitive. If the exact color is not
   confirmed by the source text, do not present a specific color
   name as a factual product specification based only on the
   photograph. You may mention a cautious visual impression in
   visual_observations, but do not use it as a confirmed product fact.

7. If a visual characteristic is uncertain, report it in warnings
   instead of presenting it as a confirmed fact.

8. Do not mention Etsy in the generated public-facing description.

9. Do not use keyword stuffing.

10. Write natural, professional English for real customers.

11. The product description should contain useful concrete
    information rather than generic marketing language.

12. Use the provided dimensions and measurements exactly as given.

Generate:

- visual_observations
- product_details
- long_description
- seo_title
- meta_description
- warnings

13. Do not mention or describe the background, table, floor,
    outdoor location, lighting setup or other photographic context.
    Only describe visible characteristics of the product itself.

14. Preserve all numerical values exactly, but normalize the
    formatting of units, spacing and multiplication symbols.

15. Create product_details as a concise list of confirmed product specifications.
    Use only facts explicitly supported by the provided source text.
    Do not infer specifications from the image.
    Do not invent or estimate missing specifications.
    Include useful measurable or functional details when available,
    such as material, dimensions, thickness, width, capacity,
    closure, hardware, stitching, belt compatibility, strap length,
    number of pockets, holes or other confirmed construction details.
    Do not include generic marketing claims.    
"""


# ---------------------------------------------------------
# 6. Structured output schema
# ---------------------------------------------------------

SCHEMA = {
    "type": "object",

    "properties": {

        "visual_observations": {
            "type": "array",

            "items": {
                "type": "string"
            }
        },
        "product_details": {
            "type": "array",

            "items": {
                "type": "string"
            }
        },
        "long_description": {
            "type": "string"
        },

        "seo_title": {
            "type": "string"
        },

        "meta_description": {
            "type": "string"
        },

        "warnings": {
            "type": "array",

            "items": {
                "type": "string"
            }
        }

    },

    "required": [
        "visual_observations",
        "product_details",
        "long_description",
        "seo_title",
        "meta_description",
        "warnings"
    ],

    "additionalProperties": False
}


# ---------------------------------------------------------
# 7. Generate with AI
# ---------------------------------------------------------

def generate_with_ai(
    product: dict,
    etsy_description: str,
    image_path: Path
) -> dict:

    load_dotenv()

    api_key = os.getenv(
        "OPENAI_API_KEY"
    )

    if not api_key:

        raise RuntimeError(
            "OPENAI_API_KEY was not found in .env"
        )


    image_bytes = (
        image_path.read_bytes()
    )

    image_base64 = (
        base64.b64encode(
            image_bytes
        ).decode("utf-8")
    )


    mime_types = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
    }


    mime_type = mime_types.get(
        image_path.suffix.lower()
    )


    if not mime_type:

        raise ValueError(
            f"Unsupported image type: "
            f"{image_path.suffix}"
        )


    image_data_url = (
        f"data:{mime_type};"
        f"base64,{image_base64}"
    )


    prompt = build_prompt(
        product,
        etsy_description
    )


    client = OpenAI(
        api_key=api_key
    )


    response = client.responses.create(

        model="gpt-5.6-luna",

        input=[
            {
                "role": "user",

                "content": [

                    {
                        "type": "input_text",
                        "text": prompt
                    },

                    {
                        "type": "input_image",
                        "image_url": image_data_url,
                        "detail": "high"
                    }

                ]
            }
        ],

        text={
            "format": {
                "type": "json_schema",
                "name": "product_content",
                "strict": True,
                "schema": SCHEMA
            }
        }
    )


    result_text = (
        response.output_text
    )


    try:

        result = json.loads(
            result_text
        )

    except json.JSONDecodeError as exc:

        raise RuntimeError(
            "OpenAI returned invalid JSON:\n"
            + result_text
        ) from exc


    return result


# ---------------------------------------------------------
# 8. Read ACTIVE products from Excel
# ---------------------------------------------------------

def load_active_products() -> list[dict]:

    if not EXCEL_PATH.exists():

        raise FileNotFoundError(
            f"Excel file not found: "
            f"{EXCEL_PATH}"
        )


    workbook = load_workbook(

        EXCEL_PATH,

        read_only=True,

        data_only=True

    )


    sheet = workbook.active


    headers = [

        cell.value

        for cell in sheet[1]

    ]


    required_headers = [

        "Category",
        "Product",
        "Short Description",
        "Etsy URL",
        "Image",
        "Alt Text",
        "Status",
        "Featured",
        "Price",
        "Currency",

    ]


    missing_headers = [

        header

        for header
        in required_headers

        if header
        not in headers

    ]


    if missing_headers:

        raise RuntimeError(

            "Missing Excel headers: "

            + ", ".join(
                missing_headers
            )

        )


    header_index = {

        header: index

        for index,
        header

        in enumerate(
            headers
        )

    }


    products = []


    for row in sheet.iter_rows(

        min_row=2,

        values_only=True

    ):


        status = clean_text(

            row[
                header_index[
                    "Status"
                ]
            ]

        ).upper()


        if status != "ACTIVE":

            continue


        etsy_url = clean_text(

            row[
                header_index[
                    "Etsy URL"
                ]
            ]

        )


        listing_id = extract_listing_id(
            etsy_url
        )


        if not listing_id:

            raise RuntimeError(

                "ACTIVE product has no valid "

                "listing ID: "

                + etsy_url

            )


        product = {

            "listing_id":
                listing_id,

            "category":
                clean_text(
                    row[
                        header_index[
                            "Category"
                        ]
                    ]
                ),

            "name":
                clean_text(
                    row[
                        header_index[
                            "Product"
                        ]
                    ]
                ),

            "short_description":
                clean_text(
                    row[
                        header_index[
                            "Short Description"
                        ]
                    ]
                ),

            "etsy_url":
                etsy_url,

            "image":
                clean_text(
                    row[
                        header_index[
                            "Image"
                        ]
                    ]
                ),

            "alt_text":
                clean_text(
                    row[
                        header_index[
                            "Alt Text"
                        ]
                    ]
                ),

            "status":
                status,

            "featured":
                clean_text(
                    row[
                        header_index[
                            "Featured"
                        ]
                    ]
                ),

            "price":
                row[
                    header_index[
                        "Price"
                    ]
                ],

            "currency":
                clean_text(
                    row[
                        header_index[
                            "Currency"
                        ]
                    ]
                ),

        }


        products.append(
            product
        )


    workbook.close()


    return products


# ---------------------------------------------------------
# 9. Process one product
# ---------------------------------------------------------

def process_product(
    product: dict
) -> str:


    listing_id = (
        product[
            "listing_id"
        ]
    )


    description_path = (

        PRODUCT_SOURCES_DIR

        / f"{listing_id}.txt"

    )


    image_path = Path(
        product[
            "image"
        ]
    )


    output_path = (

        AI_RESULTS_DIR

        / f"{listing_id}.json"

    )


    print()

    print("-" * 70)

    print(
        f"LISTING: "
        f"{listing_id}"
    )

    print(
        f"PRODUCT: "
        f"{product['name']}"
    )


    if not description_path.exists():

        raise FileNotFoundError(

            "Source description not found: "

            f"{description_path}"

        )


    if not image_path.exists():

        raise FileNotFoundError(

            "Product image not found: "

            f"{image_path}"

        )


    etsy_description = (
        description_path.read_text(
            encoding="utf-8"
        )
    )


    current_source_hash = (
        calculate_source_hash(

            product_name=
                product["name"],

            category=
                product["category"],

            short_description=
                product[
                    "short_description"
                ],

            alt_text=
                product["alt_text"],

            original_description=
                etsy_description,

            image_path=
                image_path

        )
    )


    cached = (
        load_cached_result(
            output_path
        )
    )


    if cached:

        cached_hash = (
            cached.get(
                "source_hash"
            )
        )


        if (

            cached_hash
            == current_source_hash

            and

            isinstance(
                cached.get(
                    "ai_result"
                ),
                dict
            )

        ):

            print(
                "STATUS: SKIP"
            )

            print(
                "Reason: "
                "source_hash unchanged."
            )

            print(
                f"JSON: "
                f"{output_path}"
            )


            return "skipped"


        print(
            "STATUS: REGENERATE"
        )

        print(
            "Reason: "
            "source_hash changed "
            "or cached JSON is invalid."
        )


    else:

        print(
            "STATUS: GENERATE"
        )

        print(
            "Reason: "
            "JSON does not exist."
        )


    # ---------------------------------------------
    # AI API
    # ---------------------------------------------

    result = (
        generate_with_ai(

            product=
                product,

            etsy_description=
                etsy_description,

            image_path=
                image_path

        )
    )


    # ---------------------------------------------
    # Save new JSON
    # ---------------------------------------------

    result_with_source = {

        "listing_id":
            listing_id,

        "product_name":
            product["name"],

        "category":
            product["category"],

        "etsy_url":
            product["etsy_url"],

        "image":
            str(
                image_path
            ),

        "description_source":
            str(
                description_path
            ),

        "source_hash":
            current_source_hash,

        "ai_result":
            result,

    }


    AI_RESULTS_DIR.mkdir(
        parents=True,
        exist_ok=True
    )


    output_path.write_text(

        json.dumps(

            result_with_source,

            ensure_ascii=False,

            indent=2

        ),

        encoding="utf-8"

    )


    print(
        "STATUS: "
        "AI RESULT SAVED"
    )

    print(
        f"JSON: "
        f"{output_path}"
    )


    return "generated"


# ---------------------------------------------------------
# 10. Main
# ---------------------------------------------------------

def main():


    print()

    print("=" * 70)

    print(
        "LEATHERBAGSKINGDOM "
        "- AI CONTENT GENERATOR v1"
    )

    print("=" * 70)


    products = (
        load_active_products()
    )


    print()

    print(
        f"ACTIVE PRODUCTS: "
        f"{len(products)}"
    )


    if not products:

        print()

        print(
            "No ACTIVE products found."
        )

        return


    skipped = 0

    generated = 0

    errors = 0


    for product in products:


        try:


            action = (
                process_product(
                    product
                )
            )


            if action == "skipped":

                skipped += 1


            elif action == "generated":

                generated += 1


        except Exception as exc:


            errors += 1


            print()

            print(
                "ERROR:"
            )

            print(
                str(exc)
            )

            print(
                "This product "
                "was NOT overwritten."
            )


    print()

    print("=" * 70)

    print(
        "AI GENERATION SUMMARY"
    )

    print("=" * 70)


    print(
        f"ACTIVE PRODUCTS:       "
        f"{len(products)}"
    )

    print(
        f"SKIPPED / HASH SAME:   "
        f"{skipped}"
    )

    print(
        f"GENERATED / CHANGED:   "
        f"{generated}"
    )

    print(
        f"ERRORS:                "
        f"{errors}"
    )

    print(
        f"AI API CALLS:          "
        f"{generated}"
    )

    print("=" * 70)


    if errors:

        print()

        print(
            "Completed with errors."
        )


    else:

        print()

        print(
            "Completed successfully."
        )


# ---------------------------------------------------------
# 11. Start
# ---------------------------------------------------------

if __name__ == "__main__":

    main()