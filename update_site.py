
from pathlib import Path
from openpyxl import load_workbook
from urllib.parse import urlparse
from html import escape
from datetime import datetime
import shutil


# ============================================================
# LeatherBagsKingdom
# update_site.py
#
# Excel -> Website
#
# Updates:
#   1. Featured Collections
#   2. Bags
#   3. Wallets
#   4. Belts
#   5. EDC
#
# SOLD products are not displayed.
# Short Description is taken from Excel.
# ============================================================


BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "goods(4).xlsx"
HTML_FILE = BASE_DIR / "index.html"

CATEGORIES = [
    "Bags",
    "Wallets",
    "Belts",
    "EDC"
]

VALID_STATUSES = {
    "ACTIVE",
    "NEW",
    "SOLD"
}

DISPLAY_STATUSES = {
    "ACTIVE",
    "NEW"
}

REQUIRED_HEADERS = [
    "Category",
    "Product",
    "Short Description",
    "Etsy URL",
    "Image",
    "Status",
    "Featured"
]


# ============================================================
# HELPERS
# ============================================================

def print_line():
    print("-" * 72)


def is_valid_url(value):
    if not isinstance(value, str) or not value.strip():
        return False

    try:
        parsed = urlparse(value)

        return (
            parsed.scheme in {"http", "https"}
            and bool(parsed.netloc)
        )

    except Exception:
        return False


def is_safe_image_path(image):
    """
    Image path must:
    - be relative
    - not contain ..
    - stay inside website folder
    """

    if not image:
        return False

    path = Path(image)

    if path.is_absolute():
        return False

    if ".." in path.parts:
        return False

    full_path = (BASE_DIR / path).resolve()

    try:
        full_path.relative_to(BASE_DIR.resolve())
    except ValueError:
        return False

    return True


def get_backup_path():
    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    return BASE_DIR / (
        f"index_backup_{timestamp}.html"
    )


# ============================================================
# READ EXCEL
# ============================================================

def read_excel():

    errors = []

    if not EXCEL_FILE.exists():

        errors.append(
            f"Excel file not found: {EXCEL_FILE.name}"
        )

        return [], errors

    try:

        workbook = load_workbook(
            EXCEL_FILE,
            data_only=True
        )

    except Exception as e:

        errors.append(
            f"Could not open Excel file: {e}"
        )

        return [], errors

    if "Лист2" in workbook.sheetnames:
        sheet = workbook["Лист2"]
    else:
        sheet = workbook[workbook.sheetnames[0]]

    headers = []

    for cell in sheet[1]:

        value = (
            ""
            if cell.value is None
            else str(cell.value).strip()
        )

        headers.append(value)

    while headers and headers[-1] == "":
        headers.pop()

    missing_headers = [
        h
        for h in REQUIRED_HEADERS
        if h not in headers
    ]

    if missing_headers:

        for header in missing_headers:

            errors.append(
                f"Missing required column: {header}"
            )

        return [], errors

    column_index = {
        name: index
        for index, name in enumerate(headers)
    }

    products = []

    for row_number in range(
        2,
        sheet.max_row + 1
    ):

        values = [
            sheet.cell(
                row_number,
                column_index[name] + 1
            ).value
            for name in headers
        ]

        if all(
            value is None
            or str(value).strip() == ""
            for value in values
        ):
            continue

        def get(name):

            value = values[column_index[name]]

            if value is None:
                return ""

            return str(value).strip()

        product = {

            "row": row_number,

            "category": get("Category"),

            "name": get("Product"),

            "short_description":
                get("Short Description"),

            "etsy_url": get("Etsy URL"),

            "image": get("Image"),

            "status": get("Status").upper(),

            "featured": get("Featured").upper()

        }

        products.append(product)

        # ----------------------------------------------------
        # VALIDATION
        # ----------------------------------------------------

        if product["category"] not in CATEGORIES:

            errors.append(
                f"Row {row_number}: "
                f"unknown category "
                f"'{product['category']}'"
            )

        if not product["name"]:

            errors.append(
                f"Row {row_number}: "
                f"Product is empty"
            )

        if not product["short_description"]:

            errors.append(
                f"Row {row_number}: "
                f"Short Description is empty"
            )

        if not is_valid_url(
            product["etsy_url"]
        ):

            errors.append(
                f"Row {row_number}: "
                f"invalid Etsy URL"
            )

        if not product["image"]:

            errors.append(
                f"Row {row_number}: "
                f"Image is empty"
            )

        elif not is_safe_image_path(
            product["image"]
        ):

            errors.append(
                f"Row {row_number}: "
                f"unsafe image path "
                f"'{product['image']}'"
            )

        if product["status"] not in VALID_STATUSES:

            errors.append(
                f"Row {row_number}: "
                f"invalid Status "
                f"'{product['status']}'"
            )

        if product["featured"] not in {
            "YES",
            "NO"
        }:

            errors.append(
                f"Row {row_number}: "
                f"Featured must be YES or NO"
            )

    return products, errors


# ============================================================
# IMAGE CHECK
# ============================================================

def validate_images(products):

    errors = []

    for product in products:

        image_path = (
            BASE_DIR /
            product["image"]
        )

        if not image_path.exists():

            errors.append(
                f"Row {product['row']}: "
                f"image not found: "
                f"{product['image']}"
            )

        elif not image_path.is_file():

            errors.append(
                f"Row {product['row']}: "
                f"image path is not a file: "
                f"{product['image']}"
            )

    return errors


# ============================================================
# FEATURED PRODUCTS
# ============================================================

def get_featured_products(products):

    errors = []

    featured_products = []

    for category in CATEGORIES:

        candidates = [

            product

            for product in products

            if (
                product["category"] == category
                and
                product["status"]
                in DISPLAY_STATUSES
                and
                product["featured"] == "YES"
            )

        ]

        if len(candidates) == 0:

            errors.append(
                f"{category}: "
                f"no Featured product"
            )

        elif len(candidates) > 1:

            errors.append(
                f"{category}: "
                f"{len(candidates)} "
                f"Featured products found"
            )

        else:

            featured_products.append(
                candidates[0]
            )

    return featured_products, errors


# ============================================================
# BUILD CARD
# ============================================================

def build_product_card(product):

    name = escape(
        product["name"]
    )

    description = escape(
        product["short_description"]
    )

    image = escape(
        product["image"],
        quote=True
    )

    url = escape(
        product["etsy_url"],
        quote=True
    )

    return f'''<div class="card">

    <img src="{image}"
         alt="{name}">

    <h3>
        {name}
    </h3>

    <p>
        {description}
    </p>

    <a class="btn"
       href="{url}"
       target="_blank">
       View Product
    </a>

</div>'''


# ============================================================
# BUILD PRODUCTS
# ============================================================

def build_products_block(products):

    if not products:

        return '''<div class="products">

    <!-- No active products -->

</div>'''

    cards = []

    for product in products:

        cards.append(
            build_product_card(product)
        )

    cards_text = "\n\n\n".join(
        cards
    )

    return f'''<div class="products">

{cards_text}

</div>'''


# ============================================================
# REPLACE MARKER BLOCK
# ============================================================

def replace_marker_block(
    html,
    start_marker,
    end_marker,
    new_content
):

    start = html.find(
        start_marker
    )

    end = html.find(
        end_marker
    )

    if start == -1:
        return None

    if end == -1:
        return None

    if end <= start:
        return None

    content_start = (
        start + len(start_marker)
    )

    new_html = (
        html[:content_start]
        + "\n\n"
        + new_content
        + "\n\n"
        + html[end:]
    )

    return new_html


# ============================================================
# MAIN
# ============================================================

def main():

    print()

    print("=" * 72)

    print(
        "LeatherBagsKingdom - update_site.py"
    )

    print(
        "EXCEL -> WEBSITE"
    )

    print("=" * 72)

    print()

    # --------------------------------------------------------
    # STEP 1
    # --------------------------------------------------------

    print("STEP 1 - FILE CHECK")

    print_line()

    print(
        f"Project folder : {BASE_DIR}"
    )

    print(
        f"Excel file     : {EXCEL_FILE.name}"
    )

    print(
        f"Website file   : {HTML_FILE.name}"
    )

    print()

    if not EXCEL_FILE.exists():

        print(
            "ERROR: Excel file not found."
        )

        print(
            "Expected: goods(4).xlsx"
        )

        return

    if not HTML_FILE.exists():

        print(
            "ERROR: index.html not found."
        )

        return

    print(
        "OK - required files found."
    )

    print()

    # --------------------------------------------------------
    # STEP 2
    # --------------------------------------------------------

    print("STEP 2 - READ EXCEL")

    print_line()

    products, errors = read_excel()

    print(
        f"Products found: {len(products)}"
    )

    print()

    if errors:

        print(
            f"ERRORS FOUND: {len(errors)}"
        )

        for error in errors:

            print(
                f"  ERROR: {error}"
            )

        print()

        print(
            "STOP: Website will NOT be changed."
        )

        return

    print(
        "Excel structure and data: OK"
    )

    print()

    # --------------------------------------------------------
    # STEP 3
    # --------------------------------------------------------

    print("STEP 3 - PRODUCT STATUS")

    print_line()

    active = [
        p for p in products
        if p["status"] == "ACTIVE"
    ]

    new = [
        p for p in products
        if p["status"] == "NEW"
    ]

    sold = [
        p for p in products
        if p["status"] == "SOLD"
    ]

    print(
        f"ACTIVE : {len(active)}"
    )

    print(
        f"NEW    : {len(new)}"
    )

    print(
        f"SOLD   : {len(sold)}"
    )

    print()

    # --------------------------------------------------------
    # STEP 4
    # --------------------------------------------------------

    print("STEP 4 - IMAGE CHECK")

    print_line()

    image_errors = validate_images(
        products
    )

    if image_errors:

        for error in image_errors:

            print(
                f"ERROR: {error}"
            )

        print()

        print(
            "STOP: Website will NOT be changed."
        )

        return

    print(
        f"OK - all {len(products)} "
        f"image files found."
    )

    print()

    # --------------------------------------------------------
    # STEP 5
    # --------------------------------------------------------

    print("STEP 5 - FEATURED PRODUCTS")

    print_line()

    featured_products, featured_errors = (
        get_featured_products(products)
    )

    if featured_errors:

        for error in featured_errors:

            print(
                f"ERROR: {error}"
            )

        print()

        print(
            "STOP: Website will NOT be changed."
        )

        return

    for product in featured_products:

        print(
            f"OK  {product['category']}:"
        )

        print(
            f"    {product['name']}"
        )

        print(
            f"    Description: "
            f"{product['short_description']}"
        )

        print()

    # --------------------------------------------------------
    # STEP 6
    # --------------------------------------------------------

    print("STEP 6 - CATEGORY PRODUCTS")

    print_line()

    category_products = {}

    for category in CATEGORIES:

        items = [

            p

            for p in products

            if (
                p["category"] == category
                and
                p["status"]
                in DISPLAY_STATUSES
            )

        ]

        category_products[
            category
        ] = items

        print(
            f"{category}: {len(items)} products"
        )

        for product in items:

            marker = ""

            if product["featured"] == "YES":
                marker = " ★ FEATURED"

            print(
                f"    [{product['row']}] "
                f"{product['name']}"
                f"{marker}"
            )

        print()

    # --------------------------------------------------------
    # STEP 7
    # --------------------------------------------------------

    print("STEP 7 - READ WEBSITE")

    print_line()

    try:

        html = HTML_FILE.read_text(
            encoding="utf-8"
        )

    except Exception as e:

        print(
            f"ERROR: Could not read "
            f"index.html: {e}"
        )

        return

    print(
        "OK - index.html loaded."
    )

    print()

    # --------------------------------------------------------
    # STEP 8
    # --------------------------------------------------------

    print(
        "STEP 8 - PREPARE WEBSITE UPDATE"
    )

    print_line()

    replacements = {

        "Featured": (
            "<!-- PRODUCTS_START -->",
            "<!-- PRODUCTS_END -->",
            build_products_block(
                featured_products
            )
        ),

        "Bags": (
            "<!-- BAGS_START -->",
            "<!-- BAGS_END -->",
            build_products_block(
                category_products["Bags"]
            )
        ),

        "Wallets": (
            "<!-- WALLETS_START -->",
            "<!-- WALLETS_END -->",
            build_products_block(
                category_products["Wallets"]
            )
        ),

        "Belts": (
            "<!-- BELTS_START -->",
            "<!-- BELTS_END -->",
            build_products_block(
                category_products["Belts"]
            )
        ),

        "EDC": (
            "<!-- EDC_START -->",
            "<!-- EDC_END -->",
            build_products_block(
                category_products["EDC"]
            )
        )

    }

    new_html = html

    for section, (
        start_marker,
        end_marker,
        content
    ) in replacements.items():

        updated = replace_marker_block(
            new_html,
            start_marker,
            end_marker,
            content
        )

        if updated is None:

            print(
                f"ERROR: Could not find "
                f"{section} markers."
            )

            print(
                f"  Start: {start_marker}"
            )

            print(
                f"  End:   {end_marker}"
            )

            print()

            print(
                "STOP: Website will NOT be changed."
            )

            return

        new_html = updated

        print(
            f"OK - {section} block prepared."
        )

    print()

    # --------------------------------------------------------
    # STEP 9
    # --------------------------------------------------------

    if new_html == html:

        print(
            "No changes required."
        )

        print(
            "Website already matches Excel."
        )

        return

    print(
        "STEP 9 - CREATE BACKUP"
    )

    print_line()

    backup_file = get_backup_path()

    try:

        shutil.copy2(
            HTML_FILE,
            backup_file
        )

    except Exception as e:

        print(
            f"ERROR: Could not create backup: {e}"
        )

        print(
            "STOP: Website will NOT be changed."
        )

        return

    print(
        "Backup created:"
    )

    print(
        f"  {backup_file.name}"
    )

    print()

    # --------------------------------------------------------
    # STEP 10
    # --------------------------------------------------------

    print(
        "STEP 10 - UPDATE WEBSITE"
    )

    print_line()

    try:

        HTML_FILE.write_text(
            new_html,
            encoding="utf-8"
        )

    except Exception as e:

        print(
            f"ERROR: Could not write "
            f"index.html: {e}"
        )

        print()

        print(
            "Original website is preserved "
            "in the backup."
        )

        return

    print(
        "SUCCESS - index.html updated."
    )

    print()

    # --------------------------------------------------------
    # FINAL REPORT
    # --------------------------------------------------------

    print("=" * 72)

    print(
        "UPDATE COMPLETE"
    )

    print("=" * 72)

    print()

    print(
        "FEATURED COLLECTIONS:"
    )

    for product in featured_products:

        print(
            f"  {product['category']}: "
            f"{product['name']}"
        )

    print()

    print(
        "CATEGORY TOTALS:"
    )

    for category in CATEGORIES:

        count = len(
            category_products[category]
        )

        print(
            f"  {category}: {count}"
        )

    print()

    print(
        f"TOTAL PRODUCTS ON WEBSITE: "
        f"{sum(len(category_products[c]) for c in CATEGORIES)}"
    )

    print()

    if sold:

        print(
            f"SOLD PRODUCTS NOT DISPLAYED: "
            f"{len(sold)}"
        )

    else:

        print(
            "SOLD PRODUCTS NOT DISPLAYED: 0"
        )

    print()

    print(
        "Backup:"
    )

    print(
        f"  {backup_file.name}"
    )

    print()

    print(
        "The website is now synchronized "
        "with Excel."
    )

    print()

    print("=" * 72)


if __name__ == "__main__":
    main()
