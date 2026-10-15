from pathlib import Path
import json
import html
from openpyxl import load_workbook


# ============================================================
# LEATHERBAGSKINGDOM
# Product Page Generator 
# Generates product pages for all ACTIVE/NEW products
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "goods(4).xlsx"
AI_RESULTS_DIR = BASE_DIR / "ai_results"
PRODUCTS_DIR = BASE_DIR / "products"

# ------------------------------------------------------------
# SHOP SETTINGS (Buy now / cart / Western Bid)
# ------------------------------------------------------------

# Адрес Cloudflare Worker (после деплоя: npx wrangler deploy)
WORKER_URL = "https://lbk-checkout.rabusten.workers.dev"

# Доставка одного товара по умолчанию, USD.
# Для отдельных товаров укажите свою сумму в колонке "Shipping" в Excel
# (например 10 для мелочёвки, 20 для сумок через плечо).
SHIPPING_DEFAULT_USD = "15.00"

# Если в заказе несколько товаров: самый дорогой по доставке товар платит полную
# доставку, каждый следующий - этот процент от своей доставки.
# 100 = доставки суммируются, 50 = каждая следующая со скидкой 50%.
SHIPPING_ADDITIONAL_PERCENT = 100

PRODUCTS_JSON = BASE_DIR / "products.json"

# ------------------------------------------------------------
# GOOGLE ANALYTICS
# ------------------------------------------------------------

GOOGLE_ANALYTICS = """<!-- Google tag (gtag.js) -->
<script async src="https://www.googletagmanager.com/gtag/js?id=G-VRRT7XLQ4K"></script>
<script>
  window.dataLayer = window.dataLayer || [];
  function gtag(){dataLayer.push(arguments);}
  gtag('js', new Date());

  gtag('config', 'G-VRRT7XLQ4K');
</script>"""

# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------

def esc(value):
    """HTML escape."""
    if value is None:
        return ""
    return html.escape(str(value), quote=True)


def load_excel_product(listing_id):
    """Find product in Excel by Etsy listing ID."""

    wb = load_workbook(EXCEL_FILE, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    required = [
        "Category",
        "Product",
        "Short Description",
        "Etsy URL",
        "Image",
        "Alt Text",
        "Status",
        "Featured",
        "Price",
        "Currency",
    ]

    missing = [col for col in required if col not in headers]

    if missing:
        raise RuntimeError(
            f"Missing required Excel columns: {', '.join(missing)}"
        )

    products = []

    for row in ws.iter_rows(min_row=2, values_only=True):

        data = dict(zip(headers, row))

        etsy_url = str(data.get("Etsy URL") or "").strip()

        if not etsy_url:
            continue

        # Extract listing ID from Etsy URL
        listing_id_from_url = etsy_url.rstrip("/").split("/")[-1]

        if listing_id_from_url != listing_id:
            continue

        products.append(data)

    if not products:
        raise RuntimeError(
            f"Product {listing_id} was not found in {EXCEL_FILE.name}"
        )

    return products[0]


def load_ai_result(listing_id):
    """Load AI JSON for product."""

    json_file = AI_RESULTS_DIR / f"{listing_id}.json"

    if not json_file.exists():
        raise RuntimeError(
            f"AI result not found: {json_file}"
        )

    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    if "ai_result" not in data:
        raise RuntimeError(
            f"Invalid AI JSON: 'ai_result' section is missing"
        )

    ai = data["ai_result"]

    required = [
        "long_description",
        "seo_title",
        "meta_description",
        "visual_observations",
        "product_details",
        "warnings",
    ]

    missing = [key for key in required if key not in ai]

    if missing:
        raise RuntimeError(
            "Invalid AI JSON. Missing: " + ", ".join(missing)
        )

    return data


def normalize_image_path(image_path):
    """
    Convert Excel/JSON Windows path:

        images\\bags\\1827745902.jpg

    to website-relative path from products/ page:

        ../images/bags/1827745902.jpg
    """

    path = str(image_path).replace("\\", "/")

    if path.startswith("../"):
        return path

    return "../" + path.lstrip("/")


def category_anchor(category):
    """Return homepage anchor."""

    mapping = {
        "Bags": "bags",
        "Wallets": "wallets",
        "Belts": "belts",
        "EDC": "edc",
    }

    return mapping.get(category, "home")


def build_description_html(long_description):
    """Convert AI paragraphs into HTML paragraphs."""

    paragraphs = [
        p.strip()
        for p in str(long_description).split("\n\n")
        if p.strip()
    ]

    result = []

    for paragraph in paragraphs:

        # Make Dimensions label bold if present
        if paragraph.startswith("Dimensions:"):
            paragraph = (
                "<strong>Dimensions:</strong>"
                + esc(paragraph[len("Dimensions:"):])
            )

            result.append(f"<p>{paragraph}</p>")

        else:
            result.append(f"<p>{esc(paragraph)}</p>")

    return "\n".join(result)


def build_product_details(product, ai):
    details = ai.get("product_details", [])

    if not isinstance(details, list):
        return ""

    return "\n".join(
        f"<li>{esc(str(item))}</li>"
        for item in details
        if str(item).strip()
    )

def build_json_ld(product, ai, image_url, page_url):
    """Build Product JSON-LD."""

    name = ai["seo_title"]
    description = ai["long_description"]

    price = product.get("Price")
    currency = product.get("Currency")

    if price is None:
        raise RuntimeError("Price is missing in Excel.")

    if not currency:
        raise RuntimeError("Currency is missing in Excel.")

    try:
        price_value = float(price)
        price_string = (
            f"{price_value:.2f}".rstrip("0").rstrip(".")
        )
    except (ValueError, TypeError):
        price_string = str(price)

    data = {
        "@context": "https://schema.org",
        "@type": "Product",
        "name": name,
        "description": description,
        "image": [
            image_url
        ],
        "url": page_url,
        "brand": {
            "@type": "Brand",
            "name": "LeatherBagsKingdom"
        },
        "offers": {
            "@type": "Offer",
            "price": price_string,
            "priceCurrency": str(currency),
            "availability": "https://schema.org/InStock",
            "url": product["Etsy URL"]
        }
    }

    return json.dumps(
        data,
        ensure_ascii=False,
        indent=4
    )


# ------------------------------------------------------------
# Page builder
# ------------------------------------------------------------

def build_shop_item(product, ai_data):
    """Return item dict for products.json, or None if the product can't be sold online."""

    try:
        price = f"{float(product.get('Price')):.2f}"
    except (TypeError, ValueError):
        return None

    if float(price) <= 0:
        return None

    if str(product.get("Currency") or "").strip().upper() != "USD":
        return None

    # Shipping for one unit: Excel column "Shipping" or the default
    try:
        shipping = f"{float(product.get('Shipping')):.2f}"
    except (TypeError, ValueError):
        shipping = f"{float(SHIPPING_DEFAULT_USD):.2f}"

    listing_id = str(ai_data["listing_id"])

    # path from the site root (products.json is read from the root), not from products/
    image = normalize_image_path(
        ai_data.get("image", product.get("Image")) or ""
    )[3:]

    name = str(
        ai_data.get("product_name", product.get("Product")) or ""
    ).strip()

    return {
        "name": name,
        "price": price,
        "shipping": shipping,
        "url": f"products/{listing_id}.html",
        "description": str(product.get("Short Description") or name).strip(),
        "image": image,
    }


SIMPLE_HEADER = """<header>
<nav class="navbar">
<div class="logo"><a href="index.html"><img src="logo2.png" alt="LeatherBagsKingdom Logo"></a></div>
<ul class="nav-links">
<li><a href="index.html#home">Home</a></li>
<li><a href="index.html#bags">Bags</a></li>
<li><a href="index.html#wallets">Wallets</a></li>
<li><a href="index.html#belts">Belts</a></li>
<li><a href="index.html#edc">EDC</a></li>
<li><a href="index.html#about">About</a></li>
<li><a href="index.html#contact">Contact</a></li>
<li><a class="etsy-btn" href="https://www.etsy.com/shop/LeatherBagsKingdom" target="_blank" rel="noopener">Shop on Etsy</a></li>
</ul>
</nav>
</header>"""

SIMPLE_FOOTER = """<footer class="footer">
<div class="footer-content">
<h3>LeatherBagsKingdom</h3>
<p>Handcrafted Leather Goods by Alex Polak</p>
<p>Made with passion in Ukraine 🇺🇦</p>
<p style="margin-top:15px;">Every product is handmade from premium Italian leather.</p>
<div class="footer-links">
<a href="https://www.etsy.com/shop/LeatherBagsKingdom" target="_blank" rel="noopener">Etsy</a>
|
<a href="https://www.instagram.com/leather_bags_kingdom?igsh=cXE2azE1MjFtaWZ4" target="_blank" rel="noopener">Instagram</a>
|
<a href="https://pin.it/7lFKWtFp" target="_blank" rel="noopener">Pinterest</a>
</div>
</div>
<p>© 2026 LeatherBagsKingdom. All Rights Reserved.</p>
</footer>"""


def build_simple_page(title, heading, body_html, extra_script=""):
    """Small site-styled page placed in the site root (thank-you / cancel)."""

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
{GOOGLE_ANALYTICS}
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<meta name="robots" content="noindex">
<title>{esc(title)}</title>
<link rel="stylesheet" href="style.css">
<link rel="icon" href="favicon.ico">
</head>
<body>

{SIMPLE_HEADER}

<main>
<section class="container" style="max-width:700px;margin:70px auto;padding:0 25px;text-align:center;">
<h1>{esc(heading)}</h1>
{body_html}
</section>
</main>

{SIMPLE_FOOTER}

{extra_script}

</body>
</html>
"""


def write_shop_files(shop_items):
    """Write products.json (single source of truth for the cart and the Worker),
    plus thank-you.html and cancel.html."""

    data = {
        "api": WORKER_URL,
        "currency": "USD",
        "shipping_default": f"{float(SHIPPING_DEFAULT_USD):.2f}",
        "shipping_additional_percent": SHIPPING_ADDITIONAL_PERCENT,
        "products": shop_items,
    }

    PRODUCTS_JSON.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    (BASE_DIR / "thank-you.html").write_text(
        build_simple_page(
            "Thank you for your order | LeatherBagsKingdom",
            "Thank you for your order!",
            "<p>We are confirming your payment and will start preparing your order.</p>"
            '<p><a class="btn" href="index.html">Back to the shop</a></p>',
            '<script>try{localStorage.removeItem("wb_cart_v1")}catch(e){}</script>'
        ),
        encoding="utf-8"
    )

    (BASE_DIR / "cancel.html").write_text(
        build_simple_page(
            "Payment canceled | LeatherBagsKingdom",
            "Payment canceled",
            "<p>Your payment was not completed. Your cart is saved, you can try again at any time.</p>"
            '<p><a class="btn" href="index.html">Back to the shop</a></p>'
        ),
        encoding="utf-8"
    )

    print()
    print(
        f"SHOP - products.json ({len(shop_items)} items), "
        "thank-you.html, cancel.html written."
    )

    if "YOUR-WORKER" in WORKER_URL:
        print("WARNING: set WORKER_URL at the top of this script.")

    print()


def build_product_page(product, ai_data):

    listing_id = str(ai_data["listing_id"])

    ai = ai_data["ai_result"]

    product_name = ai_data.get(
        "product_name",
        product.get("Product")
    )

    category = ai_data.get(
        "category",
        product.get("Category")
    )

    short_description = product.get(
        "Short Description",
        ""
    )

    etsy_url = ai_data.get(
        "etsy_url",
        product.get("Etsy URL")
    )

    image_path = ai_data.get(
        "image",
        product.get("Image")
    )

    alt_text = product.get(
        "Alt Text",
        ""
    )

    price = product.get("Price")
    currency = product.get("Currency")

    image_relative = normalize_image_path(image_path)

    page_url = (
        f"https://leatherbagskingdom.com/"
        f"products/{listing_id}.html"
    )

    image_url = (
        "https://leatherbagskingdom.com/"
        + str(image_path).replace("\\", "/").lstrip("/")
    )

    anchor = category_anchor(category)

    description_html = build_description_html(
        ai["long_description"]
    )

    details_html = build_product_details(
        product,
        ai
    )

    json_ld = build_json_ld(
        product,
        ai,
        image_url,
        page_url
    )

    price_display = f"${price} {currency}"

    shop_item = build_shop_item(product, ai_data)

    if shop_item:
        shop_css = '<link rel="stylesheet" href="../shop.css">'
        shop_js = '<script src="../shop.js" defer></script>'
        shop_button = f"""<div class="buy-now-wrap">
<a class="btn"
   href="#buy-now"
   data-buy-now="{esc(listing_id)}">
Buy now
</a>
</div>"""
    else:
        shop_css = ""
        shop_js = ""
        shop_button = ""

    html_page = f"""<!DOCTYPE html>
<html lang="en">
<head>
{GOOGLE_ANALYTICS}
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">

<title>{esc(ai["seo_title"])}</title>

<meta name="description"
      content="{esc(ai["meta_description"])}">

<link rel="canonical"
      href="{page_url}">

<!-- Open Graph -->
<meta property="og:type" content="product">
<meta property="og:title"
      content="{esc(ai["seo_title"])}">
<meta property="og:description"
      content="{esc(ai["meta_description"])}">
<meta property="og:url"
      content="{page_url}">
<meta property="og:image"
      content="{image_url}">

<!-- Twitter Card -->
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:title"
      content="{esc(ai["seo_title"])}">
<meta name="twitter:description"
      content="{esc(ai["meta_description"])}">
<meta name="twitter:image"
      content="{image_url}">

<link rel="stylesheet" href="../style.css">
{shop_css}

<style>
.product-page {{
    max-width: 1200px;
    margin: 55px auto 30px auto;
    padding: 0 25px;
}}

.product-detail {{
    display: flex;
    gap: 40px;
    align-items: flex-start;
    margin-bottom: 35px;
}}

.product-image {{
    flex: 1;
    text-align: center;
}}

.product-image img {{
    width: 100%;
    max-width: 560px;
    height: auto;
    display: block;
    margin: 0 auto;
}}

.product-info {{
    flex: 1;
    padding-top: 10px;
}}

.product-info h1 {{
    font-size: 32px;
    line-height: 1.2;
    margin-top: 0;
    margin-bottom: 20px;
}}

.product-info strong {{
    font-size: 20px;
}}

.product-description,
.product-details,
.product-cta {{
    max-width: 900px;
    margin: 0 auto 30px auto;
}}

.product-description p {{
    line-height: 1.6;
    margin-bottom: 15px;
    font-size: 16px;
}}

.product-details {{
    font-size: 16px;
    line-height: 1.7;
}}

.product-details ul {{
    line-height: 1.8;
}}

.product-cta {{
    text-align: center;
}}

@media (max-width: 768px) {{
    .product-detail {{
        flex-direction: column;
        gap: 30px;
    }}

    .product-image img {{
        max-width: 100%;
    }}

    .product-info {{
        width: 100%;
    }}
}}
</style>

<link rel="icon" type="image/png"
      sizes="32x32"
      href="../favicon-32x32.png">

<link rel="icon" type="image/png"
      sizes="16x16"
      href="../favicon-16x16.png">

<link rel="apple-touch-icon"
      href="../apple-touch-icon.png">

<link rel="icon"
      href="../favicon.ico">

<!-- PRODUCT JSON-LD START -->
<script type="application/ld+json">
{json_ld}
</script>
<!-- PRODUCT JSON-LD END -->

</head>

<body>

<header>
<nav class="navbar">

<div class="logo">
<a href="../index.html">
<img src="../logo2.png"
     alt="LeatherBagsKingdom Logo">
</a>
</div>

<ul class="nav-links">

<li>
<a href="../index.html#home">Home</a>
</li>

<li>
<a href="../index.html#bags">Bags</a>
</li>

<li>
<a href="../index.html#wallets">Wallets</a>
</li>

<li>
<a href="../index.html#belts">Belts</a>
</li>

<li>
<a href="../index.html#edc">EDC</a>
</li>

<li>
<a href="../index.html#about">About</a>
</li>

<li>
<a href="../index.html#contact">Contact</a>
</li>

<li>
<a class="etsy-btn"
   href="https://www.etsy.com/shop/LeatherBagsKingdom"
   target="_blank"
   rel="noopener">
Shop on Etsy
</a>
</li>

</ul>
</nav>
</header>


<main>

<section class="container product-page">

<p style="margin-bottom:20px;">
<a href="../index.html#{anchor}">
← Back to Leather {esc(category)}
</a>
</p>


<div class="product-detail">

<div class="product-image">

<img src="{esc(image_relative)}"
     alt="{esc(alt_text)}">

</div>


<div class="product-info">

<h1>
{esc(product_name)}
</h1>

<p>
{esc(short_description)}
</p>

<p>
<strong>
Price: {esc(price_display)}
</strong>
</p>

<a class="btn"
   href="{esc(etsy_url)}"
   target="_blank"
   rel="noopener"
   onclick="gtag('event', 'click_buy_on_etsy', {{
       listing_id: '{esc(listing_id)}',
       product_name: '{esc(product_name)}',
       button_position: 'top'
   }});">
Buy on Etsy
</a>
{shop_button}
</div>
</div>


<div class="product-description">

<h2>
Handmade Leather {esc(category.rstrip("s"))}
</h2>

{description_html}

</div>


<div class="product-details">

<h2>
Product Details
</h2>

<ul>
{details_html}
</ul>

</div>


<div class="product-cta">

<h2>
Get This Handmade Leather Product
</h2>

<p>
This handmade leather product is available through the
LeatherBagsKingdom Etsy shop.
</p>

<a class="btn"
   href="{esc(etsy_url)}"
   target="_blank"
   rel="noopener"
   onclick="gtag('event', 'click_buy_on_etsy', {{
       listing_id: '{esc(listing_id)}',
       product_name: '{esc(product_name)}',
       button_position: 'bottom'
   }});">
Buy on Etsy
</a>
</div>

</section>

</main>


<footer class="footer">

<div class="footer-content">

<h3>
LeatherBagsKingdom
</h3>

<p>
Handcrafted Leather Goods by Alex Polak
</p>

<p>
Made with passion in Ukraine 🇺🇦
</p>

<p style="margin-top:15px;">
Every product is handmade from premium Italian leather.
</p>

<div class="footer-links">

<a href="https://www.etsy.com/shop/LeatherBagsKingdom"
   target="_blank"
   rel="noopener">
Etsy
</a>
|
<a href="https://www.instagram.com/leather_bags_kingdom?igsh=cXE2azE1MjFtaWZ4"
   target="_blank"
   rel="noopener">
Instagram
</a>
|
<a href="https://pin.it/7lFKWtFp"
   target="_blank"
   rel="noopener">
Pinterest
</a>

</div>
</div>

<p>
© 2026 LeatherBagsKingdom. All Rights Reserved.
</p>

</footer>

{shop_js}

</body>
</html>
"""

    return html_page


# ------------------------------------------------------------
# Main
# ------------------------------------------------------------

def main():

    print()
    print("=" * 70)
    print("LEATHERBAGSKINGDOM - PRODUCT PAGE GENERATOR")
    print("=" * 70)

    print()
    print("STEP 1 - Reading Excel...")

    wb = load_workbook(EXCEL_FILE, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    products = []

    for row in ws.iter_rows(min_row=2, values_only=True):

        data = dict(zip(headers, row))

        etsy_url = str(data.get("Etsy URL") or "").strip()

        if not etsy_url:
            continue

        listing_id = etsy_url.rstrip("/").split("/")[-1]

        status = str(
            data.get("Status") or ""
        ).strip().upper()

        if status not in ("ACTIVE", "NEW"):
            continue

        data["_listing_id"] = listing_id

        products.append(data)

    print(
        f"OK - {len(products)} ACTIVE/NEW products found."
    )

    print()
    print("STEP 2 - Generating product pages...")
    print()

    PRODUCTS_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    shop_items = {}
    generated = 0
    skipped = 0
    errors = 0

    for product in products:

        listing_id = product["_listing_id"]

        print(
            f"[{listing_id}] "
            f"{product.get('Product', '')}"
        )

        try:

            ai_data = load_ai_result(
                listing_id
            )

            page = build_product_page(
                product,
                ai_data
            )

            output_file = (
                PRODUCTS_DIR
                / f"{listing_id}.html"
            )

            output_file.write_text(
                page,
                encoding="utf-8"
            )

            shop_item = build_shop_item(product, ai_data)

            if shop_item:
                shop_items[listing_id] = shop_item
            else:
                print(
                    "    NOTE - no Buy now button "
                    "(Price missing/invalid or currency is not USD)"
                )

            generated += 1

            print(
                f"    OK -> products/{listing_id}.html"
            )

        except RuntimeError as e:

            errors += 1

            print(
                f"    ERROR -> {e}"
            )

    print()
    print("=" * 70)
    print("GENERATION SUMMARY")
    print("=" * 70)

    print(
        f"ACTIVE/NEW products: {len(products)}"
    )

    print(
        f"Generated:           {generated}"
    )

    print(
        f"Errors:              {errors}"
    )

    print(
        f"Skipped:             {skipped}"
    )

    print("=" * 70)
    print()

    write_shop_files(shop_items)


if __name__ == "__main__":
    main()